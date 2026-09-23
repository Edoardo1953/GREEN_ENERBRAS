import openpyxl
import json
import re
import datetime
import os
import sys

def eval_excel_expr(expr, row_montant, row_tva, rate_tva=0.17):
    if expr is None:
        return 0.0
    if isinstance(expr, (int, float)):
        return float(expr)
    s = str(expr).strip()
    if not s:
        return 0.0
    if not s.startswith('='):
        try:
            return float(s.replace(',', '.'))
        except:
            return 0.0
            
    formula = s[1:].strip().upper()
    # Replace $P$5 or P5 with rate_tva
    formula = re.sub(r'\$?P\$?5\b', str(rate_tva), formula)
    
    # Replace SUM(Oxx:Oyy)
    def repl_sum_o(m):
        start = int(m.group(1))
        end = int(m.group(2))
        return '(' + str(sum(row_montant.get(i, 0.0) for i in range(start, end + 1))) + ')'
    formula = re.sub(r'SUM\(O(\d+):O(\d+)\)', repl_sum_o, formula)

    # Replace SUM(Pxx:Pyy)
    def repl_sum_p(m):
        start = int(m.group(1))
        end = int(m.group(2))
        return '(' + str(sum(row_tva.get(i, 0.0) for i in range(start, end + 1))) + ')'
    formula = re.sub(r'SUM\(P(\d+):P(\d+)\)', repl_sum_p, formula)

    # Replace Oxx
    def repl_o(m):
        idx = int(m.group(1))
        return '(' + str(row_montant.get(idx, 0.0)) + ')'
    formula = re.sub(r'O(\d+)', repl_o, formula)

    # Replace Pxx
    def repl_p(m):
        idx = int(m.group(1))
        return '(' + str(row_tva.get(idx, 0.0)) + ')'
    formula = re.sub(r'P(\d+)', repl_p, formula)

    try:
        if re.match(r'^[0-9\.\+\-\*\/\(\)\s]+$', formula):
            return float(eval(formula))
    except Exception:
        pass
    return 0.0

def extract_bank_transactions(project_dir):
    search_paths = [
        os.path.join(project_dir, 'uploads', 'CONTABILITA Green Enerbras One SCSp.xlsx'),
        os.path.join(project_dir, 'CONTABILITA Green Enerbras One SCSp.xlsx'),
        os.path.join(os.path.expanduser('~'), 'Desktop', 'CONTABILITA Green Enerbras One SCSp.xlsx'),
        os.path.join(os.path.expanduser('~'), 'OneDrive', 'Desktop', 'CONTABILITA Green Enerbras One SCSp.xlsx')
    ]

    xlsx_path = None
    for p in search_paths:
        if os.path.exists(p):
            xlsx_path = p
            break

    if not xlsx_path:
        print("File CONTABILITA non trovato.", file=sys.stderr)
        return None

    wb = openpyxl.load_workbook(xlsx_path, data_only=False)
    sheet_name = None
    for name in wb.sheetnames:
        if ('compte' in name.lower() and 'banque' in name.lower()) or 'banque' in name.lower():
            sheet_name = name
            break

    if not sheet_name:
        sheet_name = wb.sheetnames[0]

    ws = wb[sheet_name]

    # Read VAT rate from P5
    rate_tva = 0.17
    try:
        if ws['P5'].value is not None:
            rate_tva = float(ws['P5'].value)
    except:
        rate_tva = 0.17

    txs = []
    row_montant = {}
    row_tva = {}
    row_total = {}

    for r in range(7, ws.max_row + 1):
        raw_date = ws.cell(r, 3).value
        statut = ws.cell(r, 8).value
        cat = ws.cell(r, 9).value
        desc = ws.cell(r, 10).value
        ordinante = ws.cell(r, 11).value
        partner = ws.cell(r, 12).value
        beneficiario = ws.cell(r, 13).value
        m_val = ws.cell(r, 15).value
        p_val = ws.cell(r, 16).value

        # Parse Montant (Col O)
        m_num = eval_excel_expr(m_val, row_montant, row_tva, rate_tva)
        row_montant[r] = m_num

        # Parse TVA (Col P)
        p_num = eval_excel_expr(p_val, row_montant, row_tva, rate_tva)
        row_tva[r] = p_num

        # Total amount
        tot_num = m_num + p_num
        row_total[r] = tot_num

        # Format date
        formatted_date = ''
        if hasattr(raw_date, 'strftime'):
            formatted_date = raw_date.strftime('%d/%m/%Y')
        elif isinstance(raw_date, (int, float)):
            dt = datetime.datetime(1899, 12, 30) + datetime.timedelta(days=raw_date)
            formatted_date = dt.strftime('%d/%m/%Y')
        elif raw_date:
            formatted_date = str(raw_date).strip()

        if r == 7:
            cat = 'Saldo iniziale'
            formatted_date = ''
            desc = ''
            partner = ''
            tot_num = 0.0
        else:
            if not formatted_date:
                continue
            if str(statut).strip().upper() == 'NON':
                continue

        # Partner resolution: Partner -> Beneficiario -> Ordinante
        partner_final = partner if (partner and str(partner).strip()) else (
            beneficiario if (beneficiario and str(beneficiario).strip()) else (
                ordinante if (ordinante and str(ordinante).strip()) else ''
            )
        )

        txs.append({
            'date': formatted_date,
            'category': str(cat).strip() if cat else '',
            'description': str(desc).strip() if desc else '',
            'partner': str(partner_final).strip(),
            'amount': round(tot_num, 4)
        })

    return txs

if __name__ == '__main__':
    project_dir = os.path.dirname(os.path.abspath(__file__))
    if len(sys.argv) > 1:
        project_dir = sys.argv[1]
    
    txs = extract_bank_transactions(project_dir)
    if txs is not None:
        print(json.dumps(txs, ensure_ascii=False, indent=2))
